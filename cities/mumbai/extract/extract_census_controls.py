"""Extract historical census controls without double-counting nested geography.

The four source districts cover more than MMR. These are 2011 observations,
not a 2026 population or a ready-made simulation demand table.
"""
from collections import Counter
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
import city

PCA_SOURCES = ('census_mumbai_pca', 'census_suburban_pca',
               'census_thane_pca', 'census_raigarh_pca')
COUNTS = {
    'No_HH': 'households_count', 'TOT_P': 'persons_count',
    'TOT_M': 'male_persons_count', 'TOT_F': 'female_persons_count',
    'P_06': 'persons_age_0_6_count', 'M_06': 'male_age_0_6_count',
    'F_06': 'female_age_0_6_count', 'P_LIT': 'literate_persons_count',
    'TOT_WORK_P': 'workers_count', 'MAINWORK_P': 'main_workers_count',
    'MAIN_OT_P': 'main_other_workers_count', 'MARGWORK_P': 'marginal_workers_count',
    'MARG_OT_P': 'marginal_other_workers_count', 'NON_WORK_P': 'non_workers_count',
}
# Published PCA categories, not a crosswalk to full-time/part-time jobs.
# Preserve the existing total-column names for downstream readers.
WORK_GROUPS = {
    'TOT_WORK': 'workers', 'MAINWORK': 'main_workers',
    'MARGWORK': 'marginal_workers',
    'MARGWORK_3_6': 'marginal_workers_3_to_6_months',
    'MARGWORK_0_3': 'marginal_workers_under_3_months',
    'NON_WORK': 'non_workers',
}
SECTORS = {'CL': 'cultivators', 'AL': 'agricultural_labourers',
           'HH': 'household_industry_workers', 'OT': 'other_workers'}
WORK_PARTITIONS = (
    ('MAINWORK', 'MAIN', '', 'main'),
    ('MARGWORK', 'MARG', '', 'marginal'),
    ('MARGWORK_3_6', 'MARG', '_3_6', 'marginal_3_to_6_months'),
    ('MARGWORK_0_3', 'MARG', '_0_3', 'marginal_under_3_months'),
)
for _total, _prefix, _suffix, _name in WORK_PARTITIONS:
    for _sector, _label in SECTORS.items():
        WORK_GROUPS[_prefix + '_' + _sector + _suffix] = _name + '_' + _label
for _field, _label in WORK_GROUPS.items():
    for _sex, _sex_label in (('P', ''), ('M', 'male_'), ('F', 'female_')):
        COUNTS.setdefault(_field + '_' + _sex, _sex_label + _label + '_count')
for _prefix, _label in (('LIT', 'literate'), ('ILL', 'illiterate')):
    for _sex, _sex_label in (('P', ''), ('M', 'male_'), ('F', 'female_')):
        COUNTS.setdefault(_sex + '_' + _prefix, _sex_label + _label + '_persons_count')
OUTPUT_INPUTS = {
    'data/processed/observed/census_2011_leaf_controls.csv': ['data/raw/population/census_*_pca_*.xlsx'],
    'data/processed/observed/census_2011_b28_commuting.csv': [
        'data/raw/demand/census_b28_maharashtra_*.xlsx', 'data/raw/population/census_*_pca_*.xlsx'],
    'data/processed/observed/_census_evidence_audit.json': [
        'data/raw/demand/census_b28_maharashtra_*.xlsx', 'data/raw/population/census_*_pca_*.xlsx'],
}


def source(source_id, category):
    p = Path(city.path('data/raw', category, 'provenance_' + source_id + '.json'))
    record = json.loads(p.read_text(encoding='utf-8'))['files'][0]
    path = Path(city.path(record['path']))
    with path.open('rb') as stream:
        if hashlib.file_digest(stream, 'sha256').hexdigest() != record['sha256']:
            raise ValueError('Source hash mismatch: ' + source_id)
    return record, path


def count(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or value < 0 or int(value) != value:
        raise ValueError('Expected a reported nonnegative integer, got ' + repr(value))
    return int(value)


def write(name, rows):
    path = Path(city.path('data/processed/observed', name))
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]), lineterminator='\n')
        writer.writeheader()
        writer.writerows(rows)


def check_work_partitions(row):
    """Check source joint cells before they become synthesis constraints."""
    checks = 0
    def equal(parent, parts):
        nonlocal checks
        if count(row[parent]) != sum(count(row[p]) for p in parts):
            raise ValueError('PCA partition failed: ' + str(row['Name']) + ' ' + parent)
        checks += 1
    for field in WORK_GROUPS:
        equal(field + '_P', [field + '_M', field + '_F'])
    for sex in ('P', 'M', 'F'):
        equal('TOT_' + sex, ['TOT_WORK_' + sex, 'NON_WORK_' + sex])
        equal('TOT_WORK_' + sex, ['MAINWORK_' + sex, 'MARGWORK_' + sex])
        equal('MARGWORK_' + sex, ['MARGWORK_3_6_' + sex, 'MARGWORK_0_3_' + sex])
        for total, prefix, suffix, _ in WORK_PARTITIONS:
            equal(total + '_' + sex, [prefix + '_' + s + suffix + '_' + sex for s in SECTORS])
        for sector in SECTORS:
            equal('MARG_' + sector + '_' + sex,
                  ['MARG_' + sector + '_3_6_' + sex, 'MARG_' + sector + '_0_3_' + sex])
        equal('TOT_' + sex, [sex + '_LIT', sex + '_ILL'])
    for prefix in ('LIT', 'ILL'):
        equal('P_' + prefix, ['M_' + prefix, 'F_' + prefix])
    return checks


def main():
    controls, audits, district_codes = [], [], set()
    for source_id in PCA_SOURCES:
        record, path = source(source_id, 'population')
        book = load_workbook(path, read_only=True, data_only=True)
        iterator = book.active.iter_rows(values_only=True)
        columns = list(next(iterator))
        rows = [dict(zip(columns, row)) for row in iterator if row[0]]
        districts = [r for r in rows if r['Level'] == 'DISTRICT' and r['TRU'] == 'Total']
        if len(districts) != 1:
            raise ValueError('Expected one district total: ' + source_id)
        district = districts[0]
        partition_checks = check_work_partitions(district)
        district_codes.add(str(district['District']))
        def town_key(row):
            return tuple(str(row[k]) for k in ('State', 'District', 'Subdistt', 'Town/Village'))
        ward_towns = {town_key(r) for r in rows if r['Level'] == 'WARD'}
        leaves = [r for r in rows if r['Level'] in ('VILLAGE', 'WARD') or
                  (r['Level'] == 'TOWN' and town_key(r) not in ward_towns)]
        reconciliation = {}
        for field, label in COUNTS.items():
            total = sum(count(row[field]) for row in leaves)
            expected = count(district[field])
            reconciliation[label] = dict(leaf_total=total, published_district_total=expected)
            if total != expected:
                raise ValueError('Leaf/district reconciliation failed: ' + source_id + ' ' + field)
        for row in leaves:
            partition_checks += check_work_partitions(row)
            if count(row['TOT_M']) + count(row['TOT_F']) != count(row['TOT_P']):
                raise ValueError('Sex counts do not reconcile: ' + row['Name'])
            entry = dict(
                geography_id=':'.join(str(row[k]) for k in
                                      ('State', 'District', 'Subdistt', 'Town/Village', 'Ward', 'EB')),
                state_code=str(row['State']), district_code=str(row['District']),
                subdistrict_code=str(row['Subdistt']), town_village_code=str(row['Town/Village']),
                ward_code=str(row['Ward']), level=row['Level'], name=row['Name'],
                rural_urban=row['TRU'], observation_year=2011,
            )
            entry.update({label: count(row[field]) for field, label in COUNTS.items()})
            entry.update(source='observed', source_id=source_id, source_sha256=record['sha256'],
                         spatial_status='Historical source district; MMR inclusion and geometry unresolved')
            controls.append(entry)
        audits.append(dict(source_id=source_id, district_name=district['Name'],
                           source_levels=dict(sorted(Counter(r['Level'] for r in rows).items())),
                           leaf_rows=len(leaves), joint_partition_checks=partition_checks,
                           reconciliation=reconciliation))
        book.close()
    ids = [r['geography_id'] for r in controls]
    if len(ids) != len(set(ids)):
        raise ValueError('Duplicate census leaf geography identifier')
    record, path = source('census_b28_maharashtra', 'demand')
    book = load_workbook(path, read_only=True, data_only=True)
    raw = list(book.active.iter_rows(values_only=True))
    distances = [(6, 'Total')] + [(i, str(raw[2][i]).strip()) for i in range(9, 36, 3)]
    commuting = []
    for row in raw[6:]:
        if row[0] != 'B0128' or str(row[2]) not in district_codes:
            continue
        for i, distance in distances:
            persons, male, female = (count(row[j]) for j in range(i, i + 3))
            if male + female != persons:
                raise ValueError('B28 sex reconciliation failed')
            commuting.append(dict(
                state_code=str(row[1]), district_code=str(row[2]), residence=row[3],
                district_name=row[4], mode_raw=row[5], distance_band_km_raw=distance,
                persons_count=persons, male_persons_count=male, female_persons_count=female,
                observation_year=2011, universe='Other workers; residence-to-work travel only',
                source='observed', source_id='census_b28_maharashtra', source_sha256=record['sha256'],
                calibration_eligible='false',
            ))
    book.close()
    write('census_2011_leaf_controls.csv', controls)
    write('census_2011_b28_commuting.csv', commuting)
    audit = dict(
        pca_district_audits=audits, leaf_control_rows=len(controls), commuting_rows=len(commuting),
        count_columns=len(COUNTS),
        source_column_mapping=COUNTS,
        joint_partition_checks=sum(a['joint_partition_checks'] for a in audits),
        limitations=[
            'Four historical source districts are not the Mumbai Metropolitan Region boundary.',
            'Thane 2011 includes areas that later became Palghar district.',
            'Census ward identifiers do not directly equal current administrative or electoral wards.',
            'Leaf counts reconcile to source district totals; no current-year projection is applied.',
            'Work duration and four worker categories are retained jointly by sex at each leaf.',
            'Main/marginal work describes annual duration, not weekly full-time/part-time hours.',
            'Worker categories are residence-based and are not workplace industry or job locations.',
            'B28 covers other workers only and combines tempo, autorickshaw and taxi in one category.',
            'Total/Rural/Urban and Total/distance-band rows overlap; do not sum those dimensions together.',
            'No current all-purpose mode share or OD matrix is inferred from commuting counts.',
        ])
    Path(city.path('data/processed/observed/_census_evidence_audit.json')).write_text(
        json.dumps(audit, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    print(json.dumps({'leaf_control_rows': len(controls), 'commuting_rows': len(commuting),
                      'districts': len(audits), 'count_columns': len(COUNTS),
                      'joint_partition_checks': audit['joint_partition_checks'],
                      'district_reconciliation': 'passed'}))


if __name__ == '__main__':
    main()
