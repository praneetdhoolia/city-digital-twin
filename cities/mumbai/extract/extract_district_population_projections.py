"""Retain the published 2011-geography district projections and audit controls."""
from collections import Counter, defaultdict
from decimal import Decimal
import csv
import json
from pathlib import Path
import re
import subprocess

from openpyxl import load_workbook
import city
from extract_census_controls import PCA_SOURCES, source, write
from extract_population_projections import audit_identity

OUTPUT_INPUTS = {
    'data/processed/observed/district_population_projections.csv': [
        'data/raw/population/iips_district_projections_2012_2031_*.pdf'],
    'data/processed/observed/district_age_population_projections.csv': [
        'data/raw/population/iips_district_projections_2012_2031_*.pdf'],
    'data/processed/acquisition/district_population_projection_audit.json': [
        'data/raw/population/iips_district_projections_2012_2031_*.pdf',
        'data/raw/population/census_*_pca_*.xlsx',
        'data/processed/observed/state_population_projections.csv',
        'data/processed/observed/bmc_population_estimates.csv'],
}
SID = 'iips_district_projections_2012_2031'
STATE = 'Maharashtra'
AGE_LABELS = {'All ages', '80+'} | {str(a) for a in range(15)} | {
    str(a) + '-' + str(a+4) for a in range(15, 80, 5)}


def evidence(record, page, name, local_code):
    return dict(source='modelled', source_id=SID, source_sha256=record['sha256'],
                source_pdf_page=page, reference_date_basis='01 March', publisher_citation_year=2022,
                state_name=STATE, district_name_as_printed=name,
                district_number_within_report_state=local_code, boundary_year=2011,
                status='published_district_projection_not_observed_current_population')


def detailed_tables(pages, record):
    rows = []
    for page_no, page in enumerate(pages, 1):
        if 'State: ' + STATE + ' (27)' not in page:
            continue
        districts = set(re.findall(r'District:\s*(.+?)\s*\((\d+)\)', page))
        if len(districts) != 1:
            raise ValueError('Ambiguous detailed projection district header')
        name, code = districts.pop()
        years = None
        for line in page.splitlines():
            header = re.fullmatch(r'\s*Single\s+((?:20\d{2}\s*){5})', line)
            if header:
                years = [int(v) for v in header[1].split()]
                continue
            match = re.fullmatch(r'\s*(All ages|\d+-\d+|\d+\+|\d+)\s+((?:\d+\s+){9}\d+)\s*', line)
            if not match:
                continue
            if years is None or match[1] not in AGE_LABELS:
                raise ValueError('Unscoped or unsupported detailed projection row')
            values = list(map(int, match[2].split()))
            for i, year in enumerate(years):
                rows.append(dict(**evidence(record, page_no, name, code), reference_year=year,
                                 age_label=match[1], male_persons_count=values[i*2],
                                 female_persons_count=values[i*2+1]))
    by_district_year = defaultdict(list)
    for row in rows:
        by_district_year[row['district_number_within_report_state'], row['reference_year']].append(row)
    codes = {code for code, _ in by_district_year}
    if codes != {str(c).zfill(2) for c in range(1, 36)}:
        raise ValueError('Detailed projection district coverage changed')
    for code in sorted(codes):
        if {year for c, year in by_district_year if c == code} != set(range(2012, 2032)):
            raise ValueError('Detailed projection year coverage changed')
        for year in range(2012, 2032):
            labels = [r['age_label'] for r in by_district_year[code, year]]
            if set(labels) != AGE_LABELS or len(labels) != len(AGE_LABELS):
                raise ValueError('Missing or duplicate age category')
    return rows, by_district_year


def summary_tables(pages, record):
    rows, state = [], None
    for page_no, page in enumerate(pages, 1):
        if not re.match(r'\s*Table\s+8\b.*Projected district wise', page):
            continue
        header = re.search(r'Year\s+((?:20\d{2}\s*){5})', page)
        for line in page.splitlines():
            heading = re.fullmatch(r'\s*\d+\s*-\s*(.+?)\s*', line)
            if heading:
                state = heading[1]
                continue
            if state != STATE:
                continue
            if not header or list(map(int, header[1].split())) != list(range(2011, 2032, 5)):
                raise ValueError('Summary year headers changed')
            years = list(map(int, header[1].split()))
            match = re.fullmatch(r'\s*\d+\s+(.+?)\s*\((\d+)\)\s+((?:[\d,]+\s+){9}[\d,]+)\s*', line)
            if not match:
                continue
            name, code, cells = match.groups()
            values = [int(v.replace(',', '')) for v in cells.split()]
            for i, year in enumerate(years):
                rows.append(dict(**evidence(record, page_no, name, code), reference_year=year,
                                 male_persons_count=values[i*2], female_persons_count=values[i*2+1]))
    keys = {(r['district_number_within_report_state'], r['reference_year']) for r in rows}
    expected = {(str(c).zfill(2), y) for c in range(1, 36) for y in range(2011, 2032, 5)}
    if keys != expected or len(rows) != len(expected):
        raise ValueError('Summary district/year coverage changed')
    return rows


def main():
    record, path = source(SID, 'population')
    text = subprocess.run(['pdftotext', '-layout', str(path), '-'], check=True,
                          capture_output=True, encoding='utf-8').stdout
    pages = text.split('\f')
    if not re.search(r'MAHARASHTRA\s+Ratio Method\s*[-–]\s*1', pages[37]):
        raise ValueError('Published projection method changed')
    detailed, groups = detailed_tables(pages, record)
    summary = summary_tables(pages, record)
    totals = {(r['district_number_within_report_state'], r['reference_year']): r
              for r in detailed if r['age_label'] == 'All ages'}
    checks, crosswalk = [], []
    for row in summary:
        key = (row['district_number_within_report_state'], row['reference_year'])
        if key in totals:
            for sex in ('male', 'female'):
                delta = row[sex + '_persons_count'] - totals[key][sex + '_persons_count']
                checks.append(dict(check='summary_vs_detailed', district=row['district_name_as_printed'],
                                   reference_year=row['reference_year'], sex=sex, delta_persons=delta,
                                   status='exact' if delta == 0 else 'source_disagreement'))
    for key, group in sorted(groups.items()):
        for sex in ('male', 'female'):
            audit_identity(checks, 'age_partition_sum', dict(district_number=key[0], reference_year=key[1], sex=sex),
                           Decimal(totals[key][sex + '_persons_count']),
                           [Decimal(r[sex + '_persons_count']) for r in group if r['age_label'] != 'All ages'],
                           quantum=Decimal('1'))
    baseline = {r['district_name_as_printed']: r for r in summary if r['reference_year'] == 2011}
    for sid in PCA_SOURCES:
        pca_record, pca_path = source(sid, 'population')
        book = load_workbook(pca_path, read_only=True, data_only=True)
        try:
            iterator = book.active.iter_rows(values_only=True)
            columns = list(next(iterator))
            district_rows = [r for row in iterator if row[0]
                             if (r := dict(zip(columns, row)))['Level'] == 'DISTRICT' and r['TRU'] == 'Total']
        finally:
            book.close()
        if len(district_rows) != 1:
            raise ValueError('Ambiguous PCA district')
        pca = district_rows[0]
        if pca['Name'] not in baseline:
            raise ValueError('No exact IIPS district name candidate for PCA source')
        published = baseline[pca['Name']]
        for sex, column in (('male', 'TOT_M'), ('female', 'TOT_F')):
            if published[sex + '_persons_count'] != pca[column]:
                raise ValueError('Projection baseline differs from census district total')
        crosswalk.append(dict(district_name=pca['Name'], census_2011_district_code=pca['District'],
                              report_district_number=published['district_number_within_report_state'],
                              pca_source_id=sid, pca_sha256=pca_record['sha256'],
                              basis='exact_district_name_and_both_2011_sex_totals',
                              current_boundary_status='2011_geography_only'))
    with Path(city.path('data/processed/observed/state_population_projections.csv')).open(encoding='utf-8', newline='') as f:
        state_rows = [r for r in csv.DictReader(f) if r['residence'] == 'Total' and r['reference_date'].endswith('-03-01')]
    for row in state_rows:
        year = int(row['reference_date'][:4])
        if year not in range(2012, 2032):
            continue
        for sex in ('male', 'female'):
            components = [r[sex + '_persons_count'] for (code, y), r in totals.items() if y == year]
            target = int(row[sex + '_thousands_count']) * 1000
            delta = sum(components) - target
            rounding_bound = Decimal('500') + Decimal(len(components)) / 2
            checks.append(dict(check='districts_vs_state_projection', reference_year=year, sex=sex,
                               delta_persons=delta, rounding_bound_persons=str(rounding_bound),
                               status='exact' if delta == 0 else 'rounding_compatible' if abs(delta) <= rounding_bound
                               else 'source_disagreement'))
    # Compare only the whole Greater Mumbai total. Ward/region boundaries are
    # not silently equated to the two census district polygons.
    municipal_comparisons = []
    district_names = {r['district_name_as_printed']: r['district_number_within_report_state'] for r in summary}
    with Path(city.path('data/processed/observed/bmc_population_estimates.csv')).open(encoding='utf-8', newline='') as f:
        for r in csv.DictReader(f):
            year = int(r['reference_year_as_printed'])
            if r['area_name'] != 'Greater Mumbai' or year not in range(2012, 2032):
                continue
            projected = sum(totals[district_names[name], year][sex + '_persons_count']
                            for name in ('Mumbai', 'Mumbai Suburban') for sex in ('male', 'female'))
            municipal_comparisons.append(dict(reference_year=year, source_id=r['source_id'],
                                              source_page_1_based=r['source_page_1_based'],
                                              published_municipal_persons=int(r['population_persons_count']),
                                              iips_two_district_projected_persons=projected,
                                              difference_persons=projected-int(r['population_persons_count']),
                                              status='different_estimates_reference_date_and_boundary_reconciliation_pending'))
    write('district_population_projections.csv', summary)
    write('district_age_population_projections.csv', detailed)
    audit = dict(source='derived', source_id=SID, source_sha256=record['sha256'],
                 districts=len({r['district_name_as_printed'] for r in summary}),
                 summary_rows=len(summary), detailed_rows=len(detailed), checks=checks,
                 check_status_counts=dict(sorted(Counter(c['status'] for c in checks).items())),
                 census_baseline_crosswalk=crosswalk, municipal_comparisons=municipal_comparisons,
                 published_method=dict(name='Ratio Method 1', source_pdf_pages=[17, 18, 38],
                                       basis='2001-2011 district share changes extended to 2021, then held constant; age shares projected analogously; childhood single ages use Sprague multipliers'),
                 limitations=['District boundaries are those of 2011. Thane includes the area subsequently separated as Palghar.',
                              'Report serial numbers and within-state district numbers are not current LGD or Census district codes.',
                              'Published projections are not independent observations of 2026 population or age composition.',
                              'A constant district share after 2021 is a publisher modelling assumption, not observed absence of migration.',
                              'No downscaling to wards, villages or households has been performed.',
                              'Differences with municipal estimates are retained; no source has been overwritten.'])
    Path(city.path('data/processed/acquisition/district_population_projection_audit.json')).write_text(
        json.dumps(audit, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    print(json.dumps({k: audit[k] for k in ('districts', 'summary_rows', 'detailed_rows', 'check_status_counts', 'municipal_comparisons')}))


if __name__ == '__main__':
    main()
