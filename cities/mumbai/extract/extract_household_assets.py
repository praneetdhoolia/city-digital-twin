"""Extract historical household percentages with explicit universes and units."""
import json
import math
from pathlib import Path

from openpyxl import load_workbook

import city
from extract_census_controls import source, write

SOURCES = ('census_hl14_thane', 'census_hl14_suburban',
           'census_hl14_city', 'census_hl14_raigad')
ASSETS = {'Bicycle': 'households_with_bicycle_pct',
          'Scooter/ Motorcycle/Moped': 'households_with_two_wheeler_pct',
          'Car/ Jeep/Van': 'households_with_car_jeep_van_pct'}
OUTPUT_INPUTS = {
    'data/processed/observed/census_2011_household_assets.csv': ['data/raw/population/census_hl14_*.xlsx'],
    'data/processed/observed/_household_assets_audit.json': ['data/raw/population/census_hl14_*.xlsx'],
}


def percentage(value):
    if value is None or value == '':
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not 0 <= value <= 100:
        raise ValueError('Unexpected percentage: ' + repr(value))
    return value


def main():
    rows, audits, seen = [], [], set()
    for source_id in SOURCES:
        record, path = source(source_id, 'population')
        book = load_workbook(path, read_only=True, data_only=True)
        sheet = book.active
        cells = list(sheet.iter_rows(values_only=True))
        heading = next(i for i,r in enumerate(cells) if r[0] == 'State Code')
        labels = [str(v).strip() if v is not None else None for v in cells[heading+1]]
        columns = {labels.index(label): name for label,name in ASSETS.items()}
        # The named group fixes the meaning of otherwise repeated labels 1-5.
        size_start = cells[heading].index('Household size')
        size_end = cells[heading].index('Ownership status')
        size_columns = list(range(size_start, size_end))
        for col in size_columns:
            columns[col] = 'household_size_' + labels[col].replace('-', '_to_').replace('+', '_plus') + '_pct'
        source_count = 0
        size_sum_deviations = []
        for row_number, row in enumerate(cells[heading+4:], start=heading+5):
            if all(value is None for value in row):
                continue
            if isinstance(row[0], bool) or not str(row[0]).isdigit():
                raise ValueError('Unrecognised data row: ' + repr(row[:10]))
            geography = [str(row[i]) for i in (0, 2, 4, 6, 7)]
            residence = row[9]
            identity = tuple(geography + [residence])
            if identity in seen:
                raise ValueError('Duplicate source geographical/residence identity: ' + repr(identity))
            seen.add(identity)
            item = dict(state_code=geography[0], district_code=geography[1],
                        subdistrict_code=geography[2], town_village_code=geography[3], ward_code=geography[4],
                        area_name=row[8], residence=residence, observation_year=2011,
                        universe='Houselisting households excluding institutional households; source geography',
                        aggregation_status='Nested geographical and residence totals retained; do not sum rows')
            item.update({name: percentage(row[col]) for col,name in columns.items()})
            values = [percentage(row[col]) for col in size_columns]
            # Published rounded percentages need not add to exactly 100.
            deviation = round(sum(values)-100, 8) if all(v is not None for v in values) else None
            size_sum_deviations.append(deviation)
            item.update(household_size_sum_minus_100_pp=deviation, source='observed',
                        source_id=source_id, source_sha256=record['sha256'],
                        source_sheet=sheet.title, source_row=row_number,
                        status='historical_control_not_current_demand')
            rows.append(item)
            source_count += 1
        audits.append(dict(source_id=source_id, source_sha256=record['sha256'], rows=source_count,
                           selected_source_columns_one_based={name:col+1 for col,name in columns.items()},
                           maximum_absolute_size_sum_deviation_pp=max(abs(x) for x in size_sum_deviations if x is not None),
                           rows_with_missing_size_percentage=sum(x is None for x in size_sum_deviations)))
        book.close()
    write('census_2011_household_assets.csv', rows)
    audit = dict(schema_version=1, source='derived', status='historical_evidence_only', rows=len(rows), sources=audits,
                 limitations=[
                     'Percentages describe households possessing an asset, not numbers of vehicles or trips.',
                     'Possession of different assets overlaps; mode percentages must not be summed to 100.',
                     'Districts, subdistricts, towns, wards and rural/urban/total rows overlap.',
                     'Houselisting denominators differ from some population-census universes; no count is inferred.',
                     'These 2011 geographical units exceed and differ from the current metropolitan region.',
                     'Rounded household-size percentages retain their published sum discrepancies.',
                     'No 2026 projection or household-level joint distribution is inferred.'])
    Path(city.path('data/processed/observed/_household_assets_audit.json')).write_text(json.dumps(audit,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(audit))


if __name__ == '__main__':
    main()
