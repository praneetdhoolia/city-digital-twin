"""Reconcile historical names as candidates; never decide model inclusion."""
from collections import Counter, defaultdict
import csv
import json
from pathlib import Path
import re

from openpyxl import load_workbook
import city
from extract_census_controls import source, write

TRANSCRIPTION = 'extract/transcriptions/scheduled_areas_1985_palghar_vasai.json'
OUTPUT_INPUTS = {
    'data/processed/observed/scheduled_area_name_candidates.csv': [
        'extract/transcriptions/scheduled_areas_1985_palghar_vasai.json',
        'data/raw/boundaries/maharashtra_scheduled_areas_order_1985_*.pdf',
        'data/processed/observed/mmr_ena_villages_2024.csv',
        'data/raw/population/census_thane_pca_*.xlsx'],
    'data/processed/acquisition/scheduled_area_name_audit.json': [
        'extract/transcriptions/scheduled_areas_1985_palghar_vasai.json',
        'data/raw/boundaries/maharashtra_scheduled_areas_order_1985_*.pdf',
        'data/processed/observed/mmr_ena_villages_2024.csv',
        'data/raw/population/census_thane_pca_*.xlsx'],
}


def normalise(name):
    # Punctuation/case equivalence only. No spelling correction or fuzzy join.
    return re.sub(r'[^a-z0-9]', '', name.casefold())


def census_candidates(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        iterator = book.active.iter_rows(values_only=True)
        columns = list(next(iterator))
        rows = [dict(zip(columns, row)) for row in iterator if row[0]]
    finally:
        book.close()
    names = {str(r['Subdistt']): r['Name'] for r in rows
             if r['Level'] == 'SUB-DISTRICT' and r['TRU'] == 'Total'}
    if not names:
        raise ValueError('Missing census subdistrict names')
    by_name = defaultdict(list)
    for row in rows:
        if row['Level'] not in ('VILLAGE', 'TOWN'):
            continue
        # A census settlement suffix is a classification, not its name.
        name = re.sub(r'\s*\((?:CT|M Cl|M Corp\.|CB|NP)\)\s*$', '', row['Name'])
        key = (normalise(names[str(row['Subdistt'])]), normalise(name))
        by_name[key].append(dict(name=row['Name'], level=row['Level'],
                                 subdistrict_code=str(row['Subdistt']),
                                 town_village_code=str(row['Town/Village'])))
    return by_name


def main():
    transcription = json.loads(Path(city.path(TRANSCRIPTION)).read_text(encoding='utf-8'))
    record, _ = source(transcription['source_id'], 'boundaries')
    if record['sha256'] != transcription['source_sha256']:
        raise ValueError('Transcription does not match the acquired gazette')
    census_record, census_path = source('census_thane_pca', 'population')
    census = census_candidates(census_path)
    ena = defaultdict(list)
    with Path(city.path('data/processed/observed/mmr_ena_villages_2024.csv')).open(
            encoding='utf-8', newline='') as stream:
        for row in csv.DictReader(stream):
            ena[(normalise(row['taluka_name']), normalise(row['village_name']))].append(row)
    rows = []
    for group in transcription['lists']:
        if len(group['names']) != group['printed_village_count']:
            raise ValueError('Transcribed count differs from the printed list heading')
        if len(set(group['names'])) != len(group['names']):
            raise ValueError('Duplicate name within a transcribed list')
        for serial, name in enumerate(group['names'], 1):
            key = (normalise(group['taluka_name']), normalise(name))
            c, e = census[key], ena[key]
            rows.append(dict(source='derived',
                             source_id=transcription['source_id'], source_sha256=record['sha256'],
                             source_pdf_page=transcription['source_pdf_page'],
                             notification_date=transcription['notification_date'],
                             district_name_as_printed=transcription['district_name_as_printed'],
                             taluka_name=group['taluka_name'], serial_number=serial,
                             village_name_as_transcribed=name,
                             transcription_status='scan_reading_uncertain' if serial in group['uncertain_serials']
                             else 'english_scan_visually_transcribed',
                             census_name_candidate_count=len(c),
                             census_name_candidates_json=json.dumps(c, sort_keys=True, ensure_ascii=False),
                             ena_name_candidate_count=len(e),
                             ena_name_candidate_serials=';'.join(x['serial_number'] for x in e),
                             ena_candidate_names=';'.join(x['village_name'] for x in e),
                             boundary_status='unresolved_no_inclusion_or_exclusion_applied'))
    write('scheduled_area_name_candidates.csv', rows)
    audit = dict(source='derived', status='name_candidates_only',
                 transcription_source_sha256=record['sha256'],
                 census_source_sha256=census_record['sha256'], rows=len(rows),
                 groups={}, limitations=[
                     'An exact normalised name is a candidate, not proof of unchanged village geography.',
                     '1985 scheduled villages, 2011 census settlements and 2024 planning areas have different dates and purposes.',
                     'No fuzzy matches, polygon edits, model exclusions or inclusion decisions were applied.',
                     'This transcription covers only the Palghar and Vasai lists, not all Maharashtra Scheduled Areas.',
                     'The effect of the 2019 notification and subsequent administrative changes requires reconciliation.',
                     'Uncertain scan readings and the other language version still require verification.'])
    for taluka in sorted({r['taluka_name'] for r in rows}):
        group = [r for r in rows if r['taluka_name'] == taluka]
        audit['groups'][taluka] = dict(
            rows=len(group), uncertain_transcriptions=sum(r['transcription_status'] == 'scan_reading_uncertain' for r in group),
            census_candidate_count_distribution=dict(sorted(Counter(r['census_name_candidate_count'] for r in group).items())),
            ena_candidate_count_distribution=dict(sorted(Counter(r['ena_name_candidate_count'] for r in group).items())),
            unmatched_census_names=[r['village_name_as_transcribed'] for r in group if not r['census_name_candidate_count']],
            shared_ena_names=[r['village_name_as_transcribed'] for r in group if r['ena_name_candidate_count']])
    Path(city.path('data/processed/acquisition/scheduled_area_name_audit.json')).write_text(
        json.dumps(audit, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    print(json.dumps({k: {x: v for x, v in group.items() if not isinstance(v, list)}
                      for k, group in audit['groups'].items()}))


if __name__ == '__main__':
    main()
